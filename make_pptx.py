# -*- coding: utf-8 -*-
"""SEM3D 기반 VPE 환경 구축 로드맵 발표자료 생성 스크립트."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- 디자인 팔레트 ----
NAVY   = RGBColor(0x1F, 0x2A, 0x44)   # 진한 남색 (제목/포인트)
BLUE   = RGBColor(0x2E, 0x6F, 0xF2)   # 강조 파랑
GRAY   = RGBColor(0x44, 0x4A, 0x54)   # 본문 텍스트
LIGHT  = RGBColor(0xEE, 0xF2, 0xFB)   # 옅은 배경
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = "맑은 고딕"

# 16:9
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, color, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if not line:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def set_text(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


# =========================================================
# 표지
# =========================================================
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, SW, SH, NAVY)
# 좌측 강조 바
add_rect(s, Inches(0.9), Inches(2.7), Inches(0.12), Inches(1.9), BLUE)

tb = s.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(11), Inches(2.2))
tf = tb.text_frame
tf.word_wrap = True
p = set_text(tf, "SEM3D 기반의 Virtual Integration을 위한", 30, WHITE, bold=True)
p2 = tf.add_paragraph()
r = p2.add_run(); r.text = "VPE 환경 구축 로드맵"
r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT

sub = s.shapes.add_textbox(Inches(1.2), Inches(4.9), Inches(11), Inches(0.6))
set_text(sub.text_frame, "Virtual Process Engineering Environment Roadmap", 15, RGBColor(0xB9, 0xC6, 0xE6))


# =========================================================
# 콘텐츠 슬라이드 헬퍼
# =========================================================
def content_slide(no, title, bullets):
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, SW, SH, WHITE)
    # 헤더 바
    add_rect(s, 0, 0, SW, Inches(1.25), NAVY)
    add_rect(s, 0, Inches(1.25), SW, Inches(0.06), BLUE)
    # 번호 원 대신 번호 박스
    numbox = s.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(0.9), Inches(0.7))
    set_text(numbox.text_frame, f"{no:02d}", 26, BLUE, bold=True)
    # 제목
    tb = s.shapes.add_textbox(Inches(1.5), Inches(0.30), Inches(11), Inches(0.7))
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tb.text_frame, title, 26, WHITE, bold=True)

    # 본문
    body = s.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(11.2), Inches(5.0))
    tf = body.text_frame
    tf.word_wrap = True
    for i, txt in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        p.line_spacing = 1.25
        r = p.add_run()
        r.text = "▪  " + txt
        r.font.size = Pt(18)
        r.font.color.rgb = GRAY
        r.font.name = FONT
    return s


# 슬라이드 1 - 배경
content_slide(
    1, "배경",
    [
        "자사에 Virtual Process Integration 환경 구축이 필요하나, 기존 SEM3D Emulator는 "
        "엔지니어가 쉽게 접근하기 어려웠다.",
        "이런 접근 장벽을 낮추고자 Smart ECO 기반의 MMDM을 통한 SEM3D Emulator 실행 방식을 구축한다.",
    ],
)

# 슬라이드 2 - 효과 (내용 대기)
s2 = content_slide(2, "효과", [])
ph = s2.shapes.add_textbox(Inches(1.1), Inches(3.0), Inches(11.0), Inches(1.0))
set_text(ph.text_frame, "(효과 내용을 입력해 주시면 채워 넣겠습니다)", 16, RGBColor(0xA0, 0xA6, 0xB0))

out = "/home/user/tcadody/SEM3D_VPE_로드맵.pptx"
prs.save(out)
print("saved:", out)
